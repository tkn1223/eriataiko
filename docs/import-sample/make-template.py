# 取り込み用.xlsx を、このフォルダの 3 枚の CSV から作り直す。
#
# 使い方（openpyxl が要る。Python の作業用の環境に入れて使う）:
#   python3 -m venv /tmp/xlsx-venv && /tmp/xlsx-venv/bin/pip install openpyxl
#   /tmp/xlsx-venv/bin/python docs/import-sample/make-template.py docs/import-sample docs/import-sample/取り込み用.xlsx
#
# CSV の列を変えたら、ここの列の幅・説明（notes）・色の条件も合わせて直すこと。
import csv, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

src, out = sys.argv[1], sys.argv[2]
def rows(name):
    with open(f"{src}/{name}.csv", encoding="utf-8") as f:
        return list(csv.reader(f))

FONT = "Arial"
base = Font(name=FONT, size=11)
head = Font(name=FONT, size=11, bold=True, color="FFFFFF")
head_fill = PatternFill("solid", fgColor="2F5597")
input_fill = PatternFill("solid", fgColor="FFF9DB")
bad_fill = PatternFill("solid", fgColor="F8CBAD")
warn_fill = PatternFill("solid", fgColor="FCE4B6")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
MAXROW = 500

def table(ws, data, widths, notes):
    for r, row in enumerate(data, start=1):
        for c, v in enumerate(row, start=1):
            if r > 1 and v.isdigit():
                v = int(v)
            cell = ws.cell(row=r, column=c, value=v if v != "" else None)
            cell.font = head if r == 1 else base
            cell.border = box
            if r == 1:
                cell.fill = head_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + c)].width = w
    for c, text in enumerate(notes, start=1):
        if text:
            ws.cell(row=1, column=c).comment = Comment(text, "見本")
    ws.freeze_panes = "A2"
    # 入力する場所は薄い黄色（見本の行も、あとから書き換えて使う）
    last_col = chr(64 + len(widths))
    ws.conditional_formatting.add(
        f"A2:{last_col}{MAXROW}",
        FormulaRule(formula=["TRUE"], fill=input_fill, stopIfTrue=False),
    )

def flag(ws, rng, formula, fill):
    # 赤・橙は、取り込みで止まりそうな書き間違いの目印（シート同士は INDIRECT で見る。Google スプレッドシートの条件付き書式は別シートを直接指せないため）
    ws.conditional_formatting.add(rng, FormulaRule(formula=[formula], fill=fill, stopIfTrue=True))

wb = Workbook()

# ---- 使い方 ----
guide = wb.active
guide.title = "使い方"
lines = [
    ("取り込み用スプレッドシート（見本）", True),
    ("", False),
    ("このファイルの「大会」「選手」「試合」の 3 つのタブを、それぞれ CSV で保存して取り込みます。", False),
    ("詳しい説明はリポジトリの docs/import.md にあります。", False),
    ("", False),
    ("■ 書く場所", True),
    ("・薄い黄色のマスに書きます。いま入っているのは架空の見本です。書き換えるか消して使ってください。", False),
    ("・見出し（青い行）は変えないでください。取り込みは見出しの名前で列を探します。", False),
    ("・見出しにカーソルを当てると、その列に何を書くかのメモが出ます。", False),
    ("", False),
    ("■ 色の意味", True),
    ("・赤：取り込むと止まりそうな書き間違い（大会タブに無いチーム・部・段、選手タブに無い選手番号、重なった選手番号、面数を超えるコートなど）", False),
    ("・橙：チームA / チームB が大会タブのチーム名ではない。決勝の「予選1位」のような空枠なら、そのままで大丈夫です", False),
    ("・色が付かなくても、取り込みのときにもう一度全部確かめます（問題があれば何も書き込まずに止まります）", False),
    ("", False),
    ("■ 運営の表から写すとき", True),
    ("・運営の対戦表（印刷用）から、このタブに数式で写すと楽です。毎年レイアウトが変わっても、直すのは数式だけで済みます", False),
    ("・コートは数字で書きます。会場の A〜I は 1〜9 に直してください（例：=CODE(\"A\")-64 で A→1）", False),
    ("・人は選手番号で書きます。番号は年をまたいでずっと同じ人に同じ番号です", False),
    ("", False),
    ("■ CSV での保存（Google スプレッドシート）", True),
    ("1. 「大会」タブを開いて、ファイル → ダウンロード → カンマ区切り形式（.csv）", False),
    ("2. 保存されたファイルの名前を「大会.csv」に変える（「ファイル名 - 大会.csv」のような名前で保存されるため）", False),
    ("3. 「選手」「試合」も同じように保存して「選手.csv」「試合.csv」にする", False),
    ("4. 3 つを 1 つのフォルダに入れて、取り込みのコマンドに渡す", False),
    ("   node --env-file=.env.local scripts/import-competition.mjs <フォルダ>", False),
    ("", False),
    ("■ 気をつけること", True),
    ("・同じ日付の大会は 1 つだけ。日付を変えると別の大会として作られます（大会名は直してよい）", False),
    ("・点が 1 つでも入ったあとは、入れ直せません（当日の記録を消さないため）", False),
    ("・このタブ（使い方）は取り込みません", False),
]
for i, (text, bold) in enumerate(lines, start=1):
    c = guide.cell(row=i, column=1, value=text or None)
    c.font = Font(name=FONT, size=14 if i == 1 else 11, bold=bold)
guide.column_dimensions["A"].width = 120

# ---- 大会 ----
ws = wb.create_sheet("大会")
table(ws, rows("大会"), [12, 40, 14], [
    "行の種類。大会・会場・面数・チーム・部・段・ゲーム数 のどれか。\n大会・会場・面数は 1 行ずつ、チーム・部・段は 1 行以上。",
    "大会：大会名\n会場：会場の名前\nチーム・部・段：その名前（書いた順が並び順）\nゲーム数：「段の名前/部の名前」（例：予選リーグ/1部）\n面数：空",
    "大会：日付（2026-11-23 の形）\n会場：その会場で使える面数\n面数：今年押さえた面数（会場の面数以下）\nチーム：チーム番号 1〜4（色）\n段：リーグ か トーナメント\nゲーム数：その段・部の最大ゲーム数\n部：空",
])
dv = DataValidation(type="list", formula1='"大会,会場,面数,チーム,部,段,ゲーム数"', allow_blank=True)
dv.error, dv.errorTitle = "大会・会場・面数・チーム・部・段・ゲーム数 から選んでください", "種類"
ws.add_data_validation(dv); dv.add(f"A2:A{MAXROW}")

# ---- 選手 ----
ws = wb.create_sheet("選手")
table(ws, rows("選手"), [12, 22, 16, 10], [
    "選手番号（整数）。年をまたいでずっと同じ人に同じ番号。重ならないこと。",
    "画面に出る名前（ニックネーム）。",
    "大会タブのチーム名。試合に出ない係の人は空。",
    "大会タブの部の名前。試合に出ない係の人は空。",
])
flag(ws, f"A2:A{MAXROW}", 'AND($A2<>"",COUNTIF($A:$A,$A2)>1)', bad_fill)
flag(ws, f"C2:C{MAXROW}", 'AND($C2<>"",COUNTIFS(INDIRECT("大会!$A:$A"),"チーム",INDIRECT("大会!$B:$B"),$C2)=0)', bad_fill)
flag(ws, f"D2:D{MAXROW}", 'AND($D2<>"",COUNTIFS(INDIRECT("大会!$A:$A"),"部",INDIRECT("大会!$B:$B"),$D2)=0)', bad_fill)

# ---- 試合 ----
ws = wb.create_sheet("試合")
table(ws, rows("試合"), [14, 18, 12, 14, 14, 8, 14, 8, 16, 8, 8, 8, 8], [
    "大会タブの段の名前（予選リーグ など）。",
    "対戦の名前（予選 第1試合、準決勝1 など）。同じ「段＋対戦」の行が 1 つの対戦（チーム対チーム）になる。",
    "対戦の並び順（1 以上の整数）。同じ対戦の行はそろえる。トーナメントは 準決勝1→準決勝2→決勝→3位決定戦 の順に 4 つ、重ならないように。",
    "A 側のチーム名。決勝で未定なら「予選1位」のような空枠の名前でよい（トーナメントの段だけ）。",
    "B 側のチーム名。書き方はチームA と同じ。",
    "大会タブの部の名前。",
    "対戦の中で何番目の試合か（1 以上の整数）。",
    "コート番号（1〜面数の整数）。会場の A〜I は 1〜9 に直して書く。",
    "そのコートで何番目の試合か（1 以上の整数）。同じコートで重ならないこと。",
    "A 側の 1 人目の選手番号。トーナメントで未定なら空でよい（リーグでは必須）。",
    "A 側の 2 人目の選手番号。",
    "B 側の 1 人目の選手番号。",
    "B 側の 2 人目の選手番号。",
])
flag(ws, f"A2:A{MAXROW}", 'AND($A2<>"",COUNTIFS(INDIRECT("大会!$A:$A"),"段",INDIRECT("大会!$B:$B"),$A2)=0)', bad_fill)
flag(ws, f"F2:F{MAXROW}", 'AND($F2<>"",COUNTIFS(INDIRECT("大会!$A:$A"),"部",INDIRECT("大会!$B:$B"),$F2)=0)', bad_fill)
flag(ws, f"D2:E{MAXROW}", 'AND(D2<>"",COUNTIFS(INDIRECT("大会!$A:$A"),"チーム",INDIRECT("大会!$B:$B"),D2)=0)', warn_fill)
flag(ws, f"H2:H{MAXROW}", 'AND($H2<>"",IFERROR($H2>INDEX(INDIRECT("大会!$C:$C"),MATCH("面数",INDIRECT("大会!$A:$A"),0)),FALSE))', bad_fill)
flag(ws, f"J2:M{MAXROW}", 'AND(J2<>"",COUNTIF(INDIRECT("選手!$A:$A"),J2)=0)', bad_fill)
for col in ("C", "G", "H", "I"):
    v = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1", allow_blank=True)
    v.error, v.errorTitle = "1 以上の整数を書いてください", "数字"
    ws.add_data_validation(v); v.add(f"{col}2:{col}{MAXROW}")

wb.save(out)
print("saved", out)
